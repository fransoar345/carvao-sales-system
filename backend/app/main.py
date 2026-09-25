import asyncio
from contextvars import ContextVar
from datetime import date, datetime, time, timedelta
from io import BytesIO
from xml.sax.saxutils import escape

from fastapi import Depends, FastAPI, HTTPException, Query, Request, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordRequestForm
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func
from sqlalchemy.orm import Session, joinedload

from . import models, schemas
from .auth import create_access_token, get_current_user, hash_password, has_permission, permission_codes, require_permission, require_roles, verify_password
from .config import get_settings
from .database import Base, engine, get_db
from .seed import seed_data
from .whatsapp import get_whatsapp_settings, notify_low_stock, notify_sale, send_daily_summary, send_whatsapp


settings = get_settings()
app = FastAPI(title=settings.app_name)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in settings.cors_origins.split(",")],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

daily_summary_sent: set[date] = set()
request_ip: ContextVar[str | None] = ContextVar("request_ip", default=None)


@app.middleware("http")
async def capture_request_context(request: Request, call_next):
    token = request_ip.set(request.client.host if request.client else None)
    try:
        return await call_next(request)
    finally:
        request_ip.reset(token)


@app.on_event("startup")
async def startup():
    Base.metadata.create_all(bind=engine)
    db = next(get_db())
    try:
        seed_data(db)
    finally:
        db.close()
    asyncio.create_task(daily_summary_worker())


async def daily_summary_worker():
    while True:
        await asyncio.sleep(60)
        db = next(get_db())
        try:
            settings_row = get_whatsapp_settings(db)
            now = datetime.now()
            should_send = settings_row.daily_summary and now.time() >= settings_row.daily_summary_time and now.date() not in daily_summary_sent
            if should_send:
                await send_daily_summary(db, now)
                daily_summary_sent.add(now.date())
        except Exception as exc:
            print(f"[DAILY SUMMARY ERROR] {exc}")
        finally:
            db.close()


def audit(db: Session, user: models.User | None, action: str, entity: str, entity_id: int | None, details: str = ""):
    log = models.AuditLog(actor_id=user.id if user else None, action=action, entity=entity, entity_id=entity_id, details=details)
    db.add(log)
    db.flush()
    db.add(models.SecurityAuditDetail(audit_log_id=log.id, profiles=", ".join(role.name for role in user.access_roles) if user else None, ip_address=request_ip.get(), module=entity, new_value=details or None))


def user_to_schema(user: models.User) -> schemas.UserOut:
    return schemas.UserOut(id=user.id, name=user.name, email=user.email, role=user.role, seller_id=user.seller_id, active=user.active, profiles=[role.name for role in user.access_roles if role.active], permissions=sorted(permission_codes(user)))


def movement_delta(kind: str, quantity: float) -> float:
    mapping = {"entrada": quantity, "saida": -quantity, "ajuste": quantity, "devolucao": quantity}
    if kind not in mapping:
        raise HTTPException(400, "Tipo de movimento invalido")
    return mapping[kind]


def sale_to_schema(sale: models.Sale) -> schemas.SaleOut:
    link = sale.customer_link
    return schemas.SaleOut(
        id=sale.id,
        seller_id=sale.seller_id,
        seller_name=sale.seller.name,
        customer_name=sale.customer_name,
        customer_phone=sale.customer_phone,
        customer_id=link.customer_id if link else None,
        customer_cnpj=link.customer.cnpj if link else None,
        price_table_name=link.price_table.name if link else None,
        occurred_at=sale.occurred_at,
        total_value=sale.total_value,
        payment_method=sale.payment_method,
        payment_due_date=sale.payment_term.due_date if sale.payment_term else None,
        delivery_address=sale.delivery_detail.delivery_address if sale.delivery_detail else (link.customer.address if link else None),
        status=sale.status,
        items=[
            schemas.SaleItemOut(
                id=item.id,
                product_id=item.product_id,
                product_name=item.product.name,
                quantity=item.quantity,
                unit_price=item.unit_price,
                subtotal=item.subtotal,
            )
            for item in sale.items
        ],
    )


def sale_load_options():
    return (
        joinedload(models.Sale.seller),
        joinedload(models.Sale.items).joinedload(models.SaleItem.product),
        joinedload(models.Sale.customer_link).joinedload(models.SaleCustomerLink.customer),
        joinedload(models.Sale.customer_link).joinedload(models.SaleCustomerLink.price_table),
        joinedload(models.Sale.payment_term),
        joinedload(models.Sale.delivery_detail),
    )


def price_table_to_schema(row: models.PriceTable) -> schemas.PriceTableOut:
    return schemas.PriceTableOut(
        id=row.id, name=row.name, description=row.description, is_default=row.is_default, active=row.active,
        items=[schemas.PriceTableItemOut(product_id=item.product_id, product_name=item.product.name, price=item.price) for item in row.items],
    )


def customer_to_schema(row: models.Customer) -> schemas.CustomerOut:
    return schemas.CustomerOut(
        id=row.id, cnpj=row.cnpj, legal_name=row.legal_name, state_registration=row.state_registration,
        address=row.address, reference_point=row.reference_point, phone=row.phone, email=row.email,
        price_table_id=row.price_table_id, price_table_name=row.price_table.name if row.price_table else None, active=row.active,
        owner_seller_id=row.ownership.seller_id if row.ownership else None,
        owner_seller_name=row.ownership.seller.name if row.ownership else None,
    )


def resolve_customer_prices(db: Session, customer_id: int | None, seller_id: int, item_ids: list[int]):
    if not customer_id:
        raise HTTPException(400, "Selecione ou cadastre um cliente")
    customer = db.query(models.Customer).options(joinedload(models.Customer.price_table), joinedload(models.Customer.ownership)).filter(models.Customer.id == customer_id).first()
    if not customer or not customer.active:
        raise HTTPException(404, "Cliente nao encontrado")
    if not customer.ownership or customer.ownership.seller_id != seller_id:
        raise HTTPException(403, "Este cliente pertence a outro vendedor")
    price_table = customer.price_table or db.query(models.PriceTable).filter(models.PriceTable.is_default.is_(True), models.PriceTable.active.is_(True)).first()
    if not price_table or not price_table.active:
        raise HTTPException(400, "Cliente sem tabela de precos ativa")
    rows = db.query(models.PriceTableItem).filter(models.PriceTableItem.price_table_id == price_table.id, models.PriceTableItem.product_id.in_(item_ids)).all()
    prices = {row.product_id: row.price for row in rows}
    missing = set(item_ids) - set(prices)
    if missing:
        product = db.get(models.Product, next(iter(missing)))
        raise HTTPException(400, f"Produto {product.name if product else ''} sem preco na tabela {price_table.name}")
    return customer, price_table, prices


def delivery_manifest_options():
    return (
        joinedload(models.DeliveryManifest.items)
        .joinedload(models.DeliveryManifestItem.sale)
        .joinedload(models.Sale.seller),
        joinedload(models.DeliveryManifest.items)
        .joinedload(models.DeliveryManifestItem.sale)
        .joinedload(models.Sale.items)
        .joinedload(models.SaleItem.product),
        joinedload(models.DeliveryManifest.items)
        .joinedload(models.DeliveryManifestItem.sale)
        .joinedload(models.Sale.customer_link)
        .joinedload(models.SaleCustomerLink.customer),
        joinedload(models.DeliveryManifest.items)
        .joinedload(models.DeliveryManifestItem.sale)
        .joinedload(models.Sale.customer_link)
        .joinedload(models.SaleCustomerLink.price_table),
    )


def get_delivery_manifest(db: Session, manifest_id: int) -> models.DeliveryManifest:
    manifest = db.query(models.DeliveryManifest).options(*delivery_manifest_options()).filter(models.DeliveryManifest.id == manifest_id).first()
    if not manifest:
        raise HTTPException(404, "Romaneio nao encontrado")
    return manifest


def delivery_manifest_to_schema(manifest: models.DeliveryManifest) -> schemas.DeliveryManifestOut:
    return schemas.DeliveryManifestOut(
        id=manifest.id,
        code=manifest.code,
        delivery_date=manifest.delivery_date,
        driver_name=manifest.driver_name,
        vehicle=manifest.vehicle,
        status=manifest.status,
        notes=manifest.notes,
        created_at=manifest.created_at,
        items=[
            schemas.DeliveryManifestItemOut(
                id=item.id,
                sale_id=item.sale_id,
                delivery_address=item.delivery_address,
                delivery_order=item.delivery_order,
                status=item.status,
                delivered_at=item.delivered_at,
                note=item.note,
                sale=sale_to_schema(item.sale),
            )
            for item in manifest.items
        ],
    )


def ensure_sale_not_in_active_manifest(db: Session, sale_id: int) -> None:
    assigned = (
        db.query(models.DeliveryManifestItem)
        .join(models.DeliveryManifest)
        .filter(models.DeliveryManifestItem.sale_id == sale_id, models.DeliveryManifest.status != "cancelado")
        .first()
    )
    if assigned:
        raise HTTPException(400, "Cancele o romaneio ativo antes de alterar esta venda")


def financial_account_balance(db: Session, account: models.FinancialAccount) -> float:
    entries = db.query(func.coalesce(func.sum(models.CashTransaction.amount), 0)).filter(models.CashTransaction.account_id == account.id, models.CashTransaction.direction == "entrada", models.CashTransaction.reversed.is_(False)).scalar() or 0
    exits = db.query(func.coalesce(func.sum(models.CashTransaction.amount), 0)).filter(models.CashTransaction.account_id == account.id, models.CashTransaction.direction == "saida", models.CashTransaction.reversed.is_(False)).scalar() or 0
    return account.initial_balance + float(entries) - float(exits)


def create_receivable(db: Session, sale: models.Sale, customer_id: int, due_date: date, user: models.User):
    receivable = models.Receivable(sale_id=sale.id, customer_id=customer_id, seller_id=sale.seller_id, due_date=due_date, original_amount=sale.total_value, status="aberto")
    db.add(receivable)
    db.flush()
    if sale.payment_method != "prazo":
        account = db.query(models.FinancialAccount).filter(models.FinancialAccount.active.is_(True)).first()
        if account:
            payment = models.ReceivablePayment(receivable_id=receivable.id, account_id=account.id, amount=sale.total_value, payment_method=sale.payment_method, responsible_id=user.id)
            db.add(payment); db.flush()
            db.add(models.CashTransaction(account_id=account.id, direction="entrada", amount=sale.total_value, source_type="receivable_payment", source_id=payment.id, description=f"Recebimento automatico da venda #{sale.id}", responsible_id=user.id))
            receivable.paid_amount = sale.total_value
            receivable.status = "pago"
    return receivable


@app.get("/api/health")
def health():
    return {"ok": True, "app": settings.app_name}


@app.post("/api/auth/login", response_model=schemas.Token)
def login(form: OAuth2PasswordRequestForm = Depends(), db: Session = Depends(get_db)):
    user = db.query(models.User).options(joinedload(models.User.access_roles).joinedload(models.AccessRole.permissions), joinedload(models.User.permission_overrides).joinedload(models.UserPermissionOverride.permission)).filter(models.User.email == form.username).first()
    if not user or not verify_password(form.password, user.password_hash):
        raise HTTPException(401, "Login ou senha invalidos")
    token = create_access_token(user)
    return {"access_token": token, "user": user_to_schema(user).model_dump()}


@app.get("/api/auth/me", response_model=schemas.UserOut)
def me(user: models.User = Depends(get_current_user)):
    return user_to_schema(user)


@app.get("/api/access/permissions")
def list_permissions(db: Session = Depends(get_db), user: models.User = Depends(require_permission("users.view"))):
    rows = db.query(models.Permission).order_by(models.Permission.module, models.Permission.name).all()
    return [{"id": row.id, "code": row.code, "module": row.module, "name": row.name, "description": row.description} for row in rows]


@app.post("/api/access/permissions")
def create_permission(payload: schemas.PermissionCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("users.change_permissions"))):
    if db.query(models.Permission).filter(models.Permission.code == payload.code).first(): raise HTTPException(400, "Permissao ja cadastrada")
    row = models.Permission(**payload.model_dump()); db.add(row); db.flush(); audit(db, user, "create", "permission", row.id, row.code); db.commit()
    return {"id": row.id, **payload.model_dump()}


@app.get("/api/access/roles")
def list_access_roles(db: Session = Depends(get_db), user: models.User = Depends(require_permission("users.view"))):
    rows = db.query(models.AccessRole).options(joinedload(models.AccessRole.permissions)).order_by(models.AccessRole.name).all()
    return [{"id": row.id, "name": row.name, "description": row.description, "active": row.active, "system": row.system, "permission_codes": sorted(permission.code for permission in row.permissions)} for row in rows]


def apply_role_payload(db: Session, row: models.AccessRole, payload):
    permissions = db.query(models.Permission).filter(models.Permission.code.in_(payload.permission_codes)).all()
    if len(permissions) != len(set(payload.permission_codes)): raise HTTPException(400, "Uma ou mais permissoes nao existem")
    row.name = payload.name.strip(); row.description = payload.description; row.permissions = permissions
    if hasattr(payload, "active"): row.active = payload.active


@app.post("/api/access/roles")
def create_access_role(payload: schemas.AccessRoleCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("users.change_permissions"))):
    if db.query(models.AccessRole).filter(func.lower(models.AccessRole.name) == payload.name.lower()).first(): raise HTTPException(400, "Perfil ja cadastrado")
    row = models.AccessRole(name=payload.name, description=payload.description); db.add(row); db.flush(); apply_role_payload(db, row, payload); audit(db, user, "create", "access_role", row.id, row.name); db.commit()
    return {"id": row.id, "name": row.name, "description": row.description, "active": row.active, "permission_codes": sorted(permission.code for permission in row.permissions)}


@app.put("/api/access/roles/{role_id}")
def update_access_role(role_id: int, payload: schemas.AccessRoleUpdate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("users.change_permissions"))):
    row = db.query(models.AccessRole).options(joinedload(models.AccessRole.permissions)).filter(models.AccessRole.id == role_id).first()
    if not row: raise HTTPException(404, "Perfil nao encontrado")
    old = f"{row.name}: {','.join(permission.code for permission in row.permissions)}"; apply_role_payload(db, row, payload)
    audit(db, user, "update", "access_role", row.id, f"Antes={old}; Depois={row.name}: {','.join(payload.permission_codes)}"); db.commit()
    return {"id": row.id, "name": row.name, "description": row.description, "active": row.active, "permission_codes": sorted(permission.code for permission in row.permissions)}


@app.get("/api/access/users")
def list_access_users(db: Session = Depends(get_db), user: models.User = Depends(require_permission("users.view"))):
    rows = db.query(models.User).options(joinedload(models.User.access_roles).joinedload(models.AccessRole.permissions), joinedload(models.User.permission_overrides).joinedload(models.UserPermissionOverride.permission)).order_by(models.User.name).all()
    return [{"id": row.id, "name": row.name, "email": row.email, "active": row.active, "role_ids": [role.id for role in row.access_roles], "profiles": [role.name for role in row.access_roles], "overrides": {override.permission.code: override.allowed for override in row.permission_overrides}, "permissions": sorted(permission_codes(row))} for row in rows]


@app.post("/api/access/users")
def create_access_user(payload: schemas.AccessUserCreate, db: Session = Depends(get_db), actor: models.User = Depends(require_permission("users.create"))):
    if db.query(models.User).filter(func.lower(models.User.email) == payload.email.lower()).first(): raise HTTPException(400, "E-mail ja cadastrado")
    roles = db.query(models.AccessRole).filter(models.AccessRole.id.in_(payload.role_ids), models.AccessRole.active.is_(True)).all()
    if len(roles) != len(set(payload.role_ids)): raise HTTPException(400, "Um ou mais perfis nao existem")
    row = models.User(name=payload.name.strip(), email=payload.email.strip().lower(), password_hash=hash_password(payload.password), role=roles[0].name.lower(), active=True)
    row.access_roles = roles; db.add(row); db.flush(); audit(db, actor, "create", "user", row.id, f"{row.email}; Perfis={','.join(role.name for role in roles)}"); db.commit()
    return {"id": row.id, "name": row.name, "email": row.email, "profiles": [role.name for role in roles]}


@app.put("/api/access/users/{user_id}")
def update_user_access(user_id: int, payload: schemas.UserAccessUpdate, db: Session = Depends(get_db), actor: models.User = Depends(require_permission("users.change_role", "users.change_permissions"))):
    row = db.query(models.User).options(joinedload(models.User.access_roles), joinedload(models.User.permission_overrides)).filter(models.User.id == user_id).first()
    if not row: raise HTTPException(404, "Usuario nao encontrado")
    roles = db.query(models.AccessRole).filter(models.AccessRole.id.in_(payload.role_ids), models.AccessRole.active.is_(True)).all()
    if len(roles) != len(set(payload.role_ids)): raise HTTPException(400, "Um ou mais perfis nao existem")
    permission_rows = {item.code: item for item in db.query(models.Permission).filter(models.Permission.code.in_(list(payload.overrides))).all()}
    if len(permission_rows) != len(payload.overrides): raise HTTPException(400, "Uma ou mais permissoes nao existem")
    old = f"Perfis={','.join(role.name for role in row.access_roles)}"; row.access_roles = roles; row.permission_overrides.clear()
    row.permission_overrides.extend(models.UserPermissionOverride(permission_id=permission_rows[code].id, allowed=allowed) for code, allowed in payload.overrides.items())
    audit(db, actor, "update_access", "user", row.id, f"Antes={old}; Depois={','.join(role.name for role in roles)}; Excecoes={payload.overrides}"); db.commit()
    return {"ok": True}


@app.get("/api/access/audit")
def list_security_audit(db: Session = Depends(get_db), user: models.User = Depends(require_permission("reports.view_audit"))):
    rows = db.query(models.AuditLog).options(joinedload(models.AuditLog.security_detail)).order_by(models.AuditLog.created_at.desc()).limit(500).all()
    return [{"id": row.id, "actor_id": row.actor_id, "profiles": row.security_detail.profiles if row.security_detail else None, "ip_address": row.security_detail.ip_address if row.security_detail else None, "module": row.entity, "action": row.action, "entity_id": row.entity_id, "details": row.details, "created_at": row.created_at} for row in rows]


@app.get("/api/products", response_model=list[schemas.ProductOut])
def list_products(db: Session = Depends(get_db), user: models.User = Depends(require_permission("products.view"))):
    return db.query(models.Product).order_by(models.Product.name).all()


@app.post("/api/products", response_model=schemas.ProductOut)
def create_product(payload: schemas.ProductCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("products.create"))):
    product = models.Product(**payload.model_dump())
    db.add(product)
    db.flush()
    default_table = db.query(models.PriceTable).filter(models.PriceTable.is_default.is_(True)).first()
    if default_table:
        default_table.items.append(models.PriceTableItem(product_id=product.id, price=product.sale_price))
    audit(db, user, "create", "product", product.id, product.name)
    db.commit()
    db.refresh(product)
    return product


@app.put("/api/products/{product_id}", response_model=schemas.ProductOut)
def update_product(product_id: int, payload: schemas.ProductUpdate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("products.edit"))):
    product = db.get(models.Product, product_id)
    if not product:
        raise HTTPException(404, "Produto nao encontrado")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(product, key, value)
    if payload.sale_price is not None:
        default_item = (
            db.query(models.PriceTableItem)
            .join(models.PriceTable)
            .filter(models.PriceTable.is_default.is_(True), models.PriceTableItem.product_id == product.id)
            .first()
        )
        if default_item:
            default_item.price = payload.sale_price
    audit(db, user, "update", "product", product.id, product.name)
    db.commit()
    db.refresh(product)
    return product


@app.delete("/api/products/{product_id}")
def delete_product(product_id: int, db: Session = Depends(get_db), user: models.User = Depends(require_permission("products.deactivate"))):
    product = db.get(models.Product, product_id)
    if not product:
        raise HTTPException(404, "Produto nao encontrado")
    product.active = False
    audit(db, user, "deactivate", "product", product.id, product.name)
    db.commit()
    return {"ok": True}


@app.get("/api/price-tables", response_model=list[schemas.PriceTableOut])
def list_price_tables(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    rows = db.query(models.PriceTable).options(joinedload(models.PriceTable.items).joinedload(models.PriceTableItem.product)).order_by(models.PriceTable.name).all()
    return [price_table_to_schema(row) for row in rows]


@app.post("/api/price-tables", response_model=schemas.PriceTableOut)
def create_price_table(payload: schemas.PriceTableCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("settings.price_tables"))):
    if db.query(models.PriceTable).filter(func.lower(models.PriceTable.name) == payload.name.strip().lower()).first():
        raise HTTPException(400, "Ja existe uma tabela com este nome")
    if len({item.product_id for item in payload.items}) != len(payload.items):
        raise HTTPException(400, "Produto repetido na tabela")
    if payload.is_default:
        db.query(models.PriceTable).update({models.PriceTable.is_default: False})
    row = models.PriceTable(name=payload.name.strip(), description=payload.description, is_default=payload.is_default, active=payload.active)
    row.items = [models.PriceTableItem(product_id=item.product_id, price=item.price) for item in payload.items]
    db.add(row); db.flush(); audit(db, user, "create", "price_table", row.id, row.name); db.commit()
    row = db.query(models.PriceTable).options(joinedload(models.PriceTable.items).joinedload(models.PriceTableItem.product)).get(row.id)
    return price_table_to_schema(row)


@app.put("/api/price-tables/{table_id}", response_model=schemas.PriceTableOut)
def update_price_table(table_id: int, payload: schemas.PriceTableUpdate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("settings.price_tables"))):
    row = db.query(models.PriceTable).options(joinedload(models.PriceTable.items)).filter(models.PriceTable.id == table_id).first()
    if not row: raise HTTPException(404, "Tabela nao encontrada")
    duplicate = db.query(models.PriceTable).filter(func.lower(models.PriceTable.name) == payload.name.strip().lower(), models.PriceTable.id != table_id).first()
    if duplicate: raise HTTPException(400, "Ja existe uma tabela com este nome")
    if len({item.product_id for item in payload.items}) != len(payload.items): raise HTTPException(400, "Produto repetido na tabela")
    if payload.is_default: db.query(models.PriceTable).filter(models.PriceTable.id != table_id).update({models.PriceTable.is_default: False})
    row.name, row.description, row.is_default, row.active = payload.name.strip(), payload.description, payload.is_default, payload.active
    existing = {item.product_id: item for item in row.items}
    incoming = {item.product_id: item.price for item in payload.items}
    for product_id, price in incoming.items():
        if product_id in existing:
            existing[product_id].price = price
        else:
            row.items.append(models.PriceTableItem(product_id=product_id, price=price))
    for product_id, item in existing.items():
        if product_id not in incoming:
            db.delete(item)
    audit(db, user, "update", "price_table", row.id, row.name); db.commit()
    row = db.query(models.PriceTable).options(joinedload(models.PriceTable.items).joinedload(models.PriceTableItem.product)).get(row.id)
    return price_table_to_schema(row)


@app.get("/api/customers", response_model=list[schemas.CustomerOut])
def list_customers(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    if not has_permission(user, "customers.view_all") and not has_permission(user, "customers.view_own"):
        raise HTTPException(403, "Permissao insuficiente")
    query = db.query(models.Customer).options(joinedload(models.Customer.price_table), joinedload(models.Customer.ownership).joinedload(models.CustomerOwnership.seller))
    if not has_permission(user, "customers.view_all"):
        query = query.join(models.CustomerOwnership).filter(models.CustomerOwnership.seller_id == user.seller_id)
    rows = query.order_by(models.Customer.legal_name).all()
    return [customer_to_schema(row) for row in rows]


@app.post("/api/customers", response_model=schemas.CustomerOut)
def create_customer(payload: schemas.CustomerCreate, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    if not has_permission(user, "customers.create"): raise HTTPException(403, "Permissao insuficiente")
    if db.query(models.Customer).filter(models.Customer.cnpj == payload.cnpj).first(): raise HTTPException(400, "CNPJ ja cadastrado")
    if payload.price_table_id and not db.get(models.PriceTable, payload.price_table_id): raise HTTPException(404, "Tabela de precos nao encontrada")
    owner_seller_id = payload.owner_seller_id if has_permission(user, "customers.transfer") else user.seller_id
    if not owner_seller_id or not db.get(models.Seller, owner_seller_id): raise HTTPException(400, "Selecione o vendedor responsavel")
    customer_data = payload.model_dump(exclude={"owner_seller_id"})
    row = models.Customer(**customer_data); db.add(row); db.flush()
    row.ownership = models.CustomerOwnership(seller_id=owner_seller_id)
    audit(db, user, "create", "customer", row.id, row.legal_name); db.commit(); db.refresh(row)
    return customer_to_schema(row)


@app.put("/api/customers/{customer_id}", response_model=schemas.CustomerOut)
def update_customer(customer_id: int, payload: schemas.CustomerUpdate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("customers.edit", "customers.transfer"))):
    row = db.get(models.Customer, customer_id)
    if not row: raise HTTPException(404, "Cliente nao encontrado")
    duplicate = db.query(models.Customer).filter(models.Customer.cnpj == payload.cnpj, models.Customer.id != customer_id).first()
    if duplicate: raise HTTPException(400, "CNPJ ja cadastrado")
    if payload.price_table_id and not db.get(models.PriceTable, payload.price_table_id): raise HTTPException(404, "Tabela de precos nao encontrada")
    if not db.get(models.Seller, payload.owner_seller_id): raise HTTPException(404, "Vendedor responsavel nao encontrado")
    for key, value in payload.model_dump(exclude={"owner_seller_id"}).items(): setattr(row, key, value)
    if row.ownership: row.ownership.seller_id = payload.owner_seller_id
    else: row.ownership = models.CustomerOwnership(seller_id=payload.owner_seller_id)
    audit(db, user, "update", "customer", row.id, row.legal_name); db.commit(); db.refresh(row)
    return customer_to_schema(row)


@app.get("/api/stock/movements", response_model=list[schemas.MovementOut])
def list_movements(product_id: int | None = None, db: Session = Depends(get_db), user: models.User = Depends(require_permission("stock.view_history"))):
    query = db.query(models.StockMovement).options(joinedload(models.StockMovement.product), joinedload(models.StockMovement.responsible))
    if product_id:
        query = query.filter(models.StockMovement.product_id == product_id)
    movements = query.order_by(models.StockMovement.occurred_at.desc()).limit(300).all()
    return [
        schemas.MovementOut(
            id=m.id,
            product_id=m.product_id,
            movement_type=m.movement_type,
            quantity=m.quantity,
            occurred_at=m.occurred_at,
            responsible_id=m.responsible_id,
            note=m.note,
            product_name=m.product.name if m.product else None,
            responsible_name=m.responsible.name if m.responsible else None,
        )
        for m in movements
    ]


@app.post("/api/stock/movements", response_model=schemas.MovementOut)
async def create_movement(payload: schemas.MovementCreate, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    movement_permissions = {"entrada": "stock.entry", "saida": "stock.exit", "ajuste": "stock.adjust", "devolucao": "stock.return"}
    if not has_permission(user, movement_permissions.get(payload.movement_type, "stock.adjust")):
        raise HTTPException(403, "Permissao insuficiente para esta movimentacao")
    product = db.get(models.Product, payload.product_id)
    if not product:
        raise HTTPException(404, "Produto nao encontrado")
    product.current_stock += movement_delta(payload.movement_type, payload.quantity)
    if product.current_stock < 0:
        raise HTTPException(400, "Estoque insuficiente")
    movement = models.StockMovement(
        product_id=payload.product_id,
        movement_type=payload.movement_type,
        quantity=payload.quantity,
        occurred_at=payload.occurred_at or datetime.utcnow(),
        responsible_id=user.id,
        note=payload.note,
    )
    db.add(movement)
    db.flush()
    audit(db, user, "create", "stock_movement", movement.id, f"{product.name}: {payload.movement_type} {payload.quantity:g}")
    db.commit()
    db.refresh(movement)
    await notify_low_stock(db, product)
    return schemas.MovementOut(
        id=movement.id,
        product_id=movement.product_id,
        movement_type=movement.movement_type,
        quantity=movement.quantity,
        occurred_at=movement.occurred_at,
        responsible_id=movement.responsible_id,
        note=movement.note,
        product_name=product.name,
        responsible_name=user.name,
    )


@app.get("/api/sellers", response_model=list[schemas.SellerOut])
def list_sellers(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    sellers = db.query(models.Seller).options(joinedload(models.Seller.user)).order_by(models.Seller.name).all()
    return [
        schemas.SellerOut(
            id=s.id,
            name=s.name,
            phone=s.phone,
            monthly_goal=s.monthly_goal,
            commission_percent=s.commission_percent,
            active=s.active,
            email=s.user.email if s.user else None,
        )
        for s in sellers
    ]


@app.post("/api/sellers", response_model=schemas.SellerOut)
def create_seller(payload: schemas.SellerCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("users.create"))):
    if db.query(models.User).filter(models.User.email == payload.email).first():
        raise HTTPException(400, "E-mail ja cadastrado")
    seller = models.Seller(
        name=payload.name,
        phone=payload.phone,
        monthly_goal=payload.monthly_goal,
        commission_percent=payload.commission_percent,
    )
    db.add(seller)
    db.flush()
    app_user = models.User(
        name=payload.name,
        email=payload.email,
        password_hash=hash_password(payload.temporary_password),
        role="vendedor",
        seller_id=seller.id,
    )
    seller_role = db.query(models.AccessRole).filter(models.AccessRole.name == "Vendedor").first()
    if seller_role:
        app_user.access_roles.append(seller_role)
    db.add(app_user)
    audit(db, user, "create", "seller", seller.id, seller.name)
    db.commit()
    return schemas.SellerOut(id=seller.id, name=seller.name, phone=seller.phone, monthly_goal=seller.monthly_goal, commission_percent=seller.commission_percent, active=seller.active, email=payload.email)


@app.put("/api/sellers/{seller_id}", response_model=schemas.SellerOut)
def update_seller(seller_id: int, payload: schemas.SellerUpdate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("users.edit"))):
    seller = db.get(models.Seller, seller_id)
    if not seller:
        raise HTTPException(404, "Vendedor nao encontrado")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(seller, key, value)
    if seller.user:
        seller.user.name = seller.name
        seller.user.active = seller.active
    audit(db, user, "update", "seller", seller.id, seller.name)
    db.commit()
    return schemas.SellerOut(id=seller.id, name=seller.name, phone=seller.phone, monthly_goal=seller.monthly_goal, commission_percent=seller.commission_percent, active=seller.active, email=seller.user.email if seller.user else None)


@app.delete("/api/sellers/{seller_id}")
def delete_seller(seller_id: int, db: Session = Depends(get_db), user: models.User = Depends(require_permission("users.deactivate"))):
    seller = db.get(models.Seller, seller_id)
    if not seller:
        raise HTTPException(404, "Vendedor nao encontrado")
    seller.active = False
    if seller.user:
        seller.user.active = False
    audit(db, user, "deactivate", "seller", seller.id, seller.name)
    db.commit()
    return {"ok": True}


@app.get("/api/sales", response_model=list[schemas.SaleOut])
def list_sales(
    start: date | None = None,
    end: date | None = None,
    seller_id: int | None = None,
    db: Session = Depends(get_db),
    user: models.User = Depends(get_current_user),
):
    query = db.query(models.Sale).options(*sale_load_options())
    if not has_permission(user, "sales.view_all") and not has_permission(user, "sales.view_own"):
        raise HTTPException(403, "Permissao insuficiente")
    if not has_permission(user, "sales.view_all"):
        query = query.filter(models.Sale.seller_id == user.seller_id)
    elif seller_id:
        query = query.filter(models.Sale.seller_id == seller_id)
    if start:
        query = query.filter(models.Sale.occurred_at >= datetime.combine(start, time.min))
    if end:
        query = query.filter(models.Sale.occurred_at <= datetime.combine(end, time.max))
    return [sale_to_schema(s) for s in query.order_by(models.Sale.occurred_at.desc()).limit(300).all()]


@app.post("/api/sales", response_model=schemas.SaleOut)
async def create_sale(payload: schemas.SaleCreate, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    if not has_permission(user, "sales.create"): raise HTTPException(403, "Permissao insuficiente")
    seller_id = user.seller_id if has_permission(user, "sales.view_own") and not has_permission(user, "sales.view_all") else payload.seller_id
    if not seller_id:
        raise HTTPException(400, "Selecione um vendedor")
    seller = db.get(models.Seller, seller_id)
    if not seller or not seller.active:
        raise HTTPException(404, "Vendedor nao encontrado")
    if not payload.items:
        raise HTTPException(400, "Inclua ao menos um item")

    customer, price_table, prices = resolve_customer_prices(db, payload.customer_id, seller_id, [item.product_id for item in payload.items])
    sale = models.Sale(seller_id=seller_id, customer_name=customer.legal_name, customer_phone=customer.phone, payment_method=payload.payment_method, status="confirmada")
    total = 0.0
    touched_products: list[models.Product] = []
    for item in payload.items:
        product = db.get(models.Product, item.product_id)
        if not product or not product.active:
            raise HTTPException(404, "Produto indisponivel")
        if product.current_stock < item.quantity:
            raise HTTPException(400, f"Estoque insuficiente para {product.name}")
        unit_price = prices[item.product_id]
        subtotal = item.quantity * unit_price
        product.current_stock -= item.quantity
        total += subtotal
        touched_products.append(product)
        sale.items.append(models.SaleItem(product_id=product.id, quantity=item.quantity, unit_price=unit_price, subtotal=subtotal))
        db.add(models.StockMovement(product_id=product.id, movement_type="saida", quantity=item.quantity, responsible_id=user.id, note="Baixa automatica por venda"))
    sale.total_value = total
    db.add(sale)
    db.flush()
    sale.customer_link = models.SaleCustomerLink(customer_id=customer.id, price_table_id=price_table.id)
    if payload.payment_method == "prazo":
        sale.payment_term = models.SalePaymentTerm(due_date=payload.payment_due_date)
    sale.delivery_detail = models.SaleDeliveryDetail(delivery_address=payload.delivery_address.strip())
    create_receivable(db, sale, customer.id, payload.payment_due_date or datetime.utcnow().date(), user)
    audit(db, user, "create", "sale", sale.id, f"Venda R$ {sale.total_value:.2f}")
    db.commit()
    sale = db.query(models.Sale).options(*sale_load_options()).get(sale.id)
    await notify_sale(db, sale)
    for product in touched_products:
        await notify_low_stock(db, product)
    return sale_to_schema(sale)


@app.put("/api/sales/{sale_id}", response_model=schemas.SaleOut)
def update_sale(sale_id: int, payload: schemas.SaleUpdate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("sales.edit"))):
    sale = db.query(models.Sale).options(*sale_load_options()).filter(models.Sale.id == sale_id).first()
    if not sale:
        raise HTTPException(404, "Venda nao encontrada")
    if sale.status != "confirmada":
        raise HTTPException(400, "Somente vendas confirmadas podem ser alteradas")
    ensure_sale_not_in_active_manifest(db, sale.id)
    seller = db.get(models.Seller, payload.seller_id)
    if not seller or not seller.active:
        raise HTTPException(404, "Vendedor nao encontrado")
    customer, price_table, prices = resolve_customer_prices(db, payload.customer_id, payload.seller_id, [item.product_id for item in payload.items])

    restored_by_product: dict[int, float] = {}
    for old_item in sale.items:
        restored_by_product[old_item.product_id] = restored_by_product.get(old_item.product_id, 0) + old_item.quantity
    requested_by_product: dict[int, float] = {}
    for new_item in payload.items:
        requested_by_product[new_item.product_id] = requested_by_product.get(new_item.product_id, 0) + new_item.quantity
    products = {product.id: product for product in db.query(models.Product).filter(models.Product.id.in_(list(requested_by_product))).all()}
    if len(products) != len(requested_by_product) or any(not product.active for product in products.values()):
        raise HTTPException(404, "Um ou mais produtos estao indisponiveis")
    for product_id, quantity in requested_by_product.items():
        available = products[product_id].current_stock + restored_by_product.get(product_id, 0)
        if available < quantity:
            raise HTTPException(400, f"Estoque insuficiente para {products[product_id].name}")

    for old_item in list(sale.items):
        old_item.product.current_stock += old_item.quantity
        db.add(models.StockMovement(product_id=old_item.product_id, movement_type="devolucao", quantity=old_item.quantity, responsible_id=user.id, note=f"Estorno para edicao da venda #{sale.id}"))
    sale.items.clear()
    total = 0.0
    for new_item in payload.items:
        product = products[new_item.product_id]
        unit_price = prices[new_item.product_id]
        subtotal = new_item.quantity * unit_price
        product.current_stock -= new_item.quantity
        total += subtotal
        sale.items.append(models.SaleItem(product_id=product.id, quantity=new_item.quantity, unit_price=unit_price, subtotal=subtotal))
        db.add(models.StockMovement(product_id=product.id, movement_type="saida", quantity=new_item.quantity, responsible_id=user.id, note=f"Baixa por edicao da venda #{sale.id}"))
    sale.seller_id = payload.seller_id
    sale.customer_name = customer.legal_name
    sale.customer_phone = customer.phone
    sale.payment_method = payload.payment_method
    if sale.delivery_detail: sale.delivery_detail.delivery_address = payload.delivery_address.strip()
    else: sale.delivery_detail = models.SaleDeliveryDetail(delivery_address=payload.delivery_address.strip())
    if payload.payment_method == "prazo":
        if sale.payment_term:
            sale.payment_term.due_date = payload.payment_due_date
        else:
            sale.payment_term = models.SalePaymentTerm(due_date=payload.payment_due_date)
    elif sale.payment_term:
        db.delete(sale.payment_term)
        sale.payment_term = None
    sale.total_value = total
    if sale.customer_link:
        sale.customer_link.customer_id = customer.id
        sale.customer_link.price_table_id = price_table.id
    else:
        sale.customer_link = models.SaleCustomerLink(customer_id=customer.id, price_table_id=price_table.id)
    receivable = db.query(models.Receivable).options(joinedload(models.Receivable.payments)).filter(models.Receivable.sale_id == sale.id).first()
    if receivable:
        receivable.customer_id = customer.id
        receivable.seller_id = sale.seller_id
        receivable.due_date = payload.payment_due_date or sale.occurred_at.date()
        receivable.original_amount = total
        active_payments = [payment for payment in receivable.payments if not payment.reversed]
        receivable.paid_amount = sum(payment.amount for payment in active_payments)
        receivable.status = "pago" if receivable.paid_amount >= total else ("parcial" if receivable.paid_amount else "aberto")
        if len(active_payments) == 1 and sale.payment_method != "prazo":
            active_payments[0].amount = total
            receivable.paid_amount = total
            receivable.status = "pago"
            db.query(models.CashTransaction).filter(models.CashTransaction.source_type == "receivable_payment", models.CashTransaction.source_id == active_payments[0].id).update({models.CashTransaction.amount: total})
    audit(db, user, "update", "sale", sale.id, f"Venda alterada para R$ {sale.total_value:.2f}")
    db.commit()
    updated = db.query(models.Sale).options(*sale_load_options()).filter(models.Sale.id == sale.id).first()
    return sale_to_schema(updated)


@app.delete("/api/sales/{sale_id}")
def delete_sale(sale_id: int, db: Session = Depends(get_db), user: models.User = Depends(require_permission("sales.cancel"))):
    sale = db.query(models.Sale).options(joinedload(models.Sale.items).joinedload(models.SaleItem.product)).filter(models.Sale.id == sale_id).first()
    if not sale:
        raise HTTPException(404, "Venda nao encontrada")
    if sale.status != "confirmada":
        raise HTTPException(400, "Venda ja cancelada")
    ensure_sale_not_in_active_manifest(db, sale.id)
    for item in sale.items:
        item.product.current_stock += item.quantity
        db.add(models.StockMovement(product_id=item.product_id, movement_type="devolucao", quantity=item.quantity, responsible_id=user.id, note=f"Cancelamento da venda #{sale.id}"))
    sale.status = "cancelada"
    receivable = db.query(models.Receivable).options(joinedload(models.Receivable.payments)).filter(models.Receivable.sale_id == sale.id).first()
    if receivable:
        receivable.status = "cancelado"
        for payment in receivable.payments:
            payment.reversed = True
            db.query(models.CashTransaction).filter(models.CashTransaction.source_type == "receivable_payment", models.CashTransaction.source_id == payment.id).update({models.CashTransaction.reversed: True})
    audit(db, user, "cancel", "sale", sale.id, f"Venda cancelada e estoque devolvido: R$ {sale.total_value:.2f}")
    db.commit()
    return {"ok": True, "status": sale.status}


@app.get("/api/delivery-manifests/pending-sales", response_model=list[schemas.SaleOut])
def list_sales_pending_delivery(db: Session = Depends(get_db), user: models.User = Depends(require_permission("manifests.create"))):
    assigned_sale_ids = (
        db.query(models.DeliveryManifestItem.sale_id)
        .join(models.DeliveryManifest)
        .filter(models.DeliveryManifest.status != "cancelado")
    )
    sales = (
        db.query(models.Sale)
        .options(*sale_load_options())
        .filter(models.Sale.status == "confirmada", ~models.Sale.id.in_(assigned_sale_ids))
        .order_by(models.Sale.occurred_at.desc())
        .all()
    )
    return [sale_to_schema(sale) for sale in sales]


@app.get("/api/delivery-manifests", response_model=list[schemas.DeliveryManifestOut])
def list_delivery_manifests(db: Session = Depends(get_db), user: models.User = Depends(require_permission("manifests.view"))):
    query = db.query(models.DeliveryManifest).options(*delivery_manifest_options())
    if not has_permission(user, "manifests.create") and has_permission(user, "manifests.confirm_delivery"):
        query = query.filter(func.lower(models.DeliveryManifest.driver_name) == user.name.lower())
    manifests = query.order_by(models.DeliveryManifest.delivery_date.desc(), models.DeliveryManifest.id.desc()).limit(100).all()
    return [delivery_manifest_to_schema(manifest) for manifest in manifests]


@app.post("/api/delivery-manifests", response_model=schemas.DeliveryManifestOut)
def create_delivery_manifest(payload: schemas.DeliveryManifestCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("manifests.create"))):
    sale_ids = [item.sale_id for item in payload.items]
    if len(sale_ids) != len(set(sale_ids)):
        raise HTTPException(400, "Uma venda nao pode aparecer duas vezes no mesmo romaneio")
    already_assigned = (
        db.query(models.DeliveryManifestItem)
        .join(models.DeliveryManifest)
        .filter(models.DeliveryManifestItem.sale_id.in_(sale_ids), models.DeliveryManifest.status != "cancelado")
        .first()
    )
    if already_assigned:
        raise HTTPException(400, f"A venda {already_assigned.sale_id} ja pertence a um romaneio ativo")
    sales = db.query(models.Sale).filter(models.Sale.id.in_(sale_ids), models.Sale.status == "confirmada").all()
    if len(sales) != len(sale_ids):
        raise HTTPException(400, "Uma ou mais vendas nao foram encontradas")

    manifest = models.DeliveryManifest(
        code=f"TMP-{datetime.utcnow():%y%m%d%H%M%S%f}",
        delivery_date=payload.delivery_date,
        driver_name=payload.driver_name,
        vehicle=payload.vehicle,
        status="preparacao",
        notes=payload.notes,
        created_by_id=user.id,
    )
    for position, item in enumerate(sorted(payload.items, key=lambda row: row.delivery_order), start=1):
        manifest.items.append(models.DeliveryManifestItem(
            sale_id=item.sale_id,
            delivery_address=item.delivery_address.strip(),
            delivery_order=position,
            status="pendente",
        ))
    db.add(manifest)
    db.flush()
    manifest.code = f"ROM-{payload.delivery_date:%Y%m%d}-{manifest.id:04d}"
    audit(db, user, "create", "delivery_manifest", manifest.id, f"{manifest.code} com {len(manifest.items)} entregas")
    db.commit()
    return delivery_manifest_to_schema(get_delivery_manifest(db, manifest.id))


@app.put("/api/delivery-manifests/{manifest_id}/status", response_model=schemas.DeliveryManifestOut)
def update_delivery_manifest_status(manifest_id: int, payload: schemas.DeliveryManifestStatusUpdate, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    needed = {"em_rota": "manifests.start_route", "concluido": "manifests.finish_route", "cancelado": "manifests.cancel"}.get(payload.status, "manifests.edit")
    if not has_permission(user, needed): raise HTTPException(403, "Permissao insuficiente")
    allowed = {"preparacao", "em_rota", "concluido", "cancelado"}
    if payload.status not in allowed:
        raise HTTPException(400, "Status de romaneio invalido")
    manifest = get_delivery_manifest(db, manifest_id)
    if not has_permission(user, "manifests.create") and manifest.driver_name.lower() != user.name.lower():
        raise HTTPException(403, "Romaneio atribuido a outro motorista")
    if payload.status == "concluido" and any(item.status != "entregue" for item in manifest.items):
        raise HTTPException(400, "Confirme todas as entregas antes de concluir o romaneio")
    manifest.status = payload.status
    if payload.status == "cancelado":
        for item in manifest.items:
            if item.status != "entregue":
                item.status = "cancelado"
    audit(db, user, "status", "delivery_manifest", manifest.id, payload.status)
    db.commit()
    return delivery_manifest_to_schema(get_delivery_manifest(db, manifest.id))


@app.put("/api/delivery-manifests/{manifest_id}/items/{item_id}", response_model=schemas.DeliveryManifestOut)
def update_delivery_item(manifest_id: int, item_id: int, payload: schemas.DeliveryItemStatusUpdate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("manifests.confirm_delivery"))):
    allowed = {"pendente", "entregue", "nao_entregue"}
    if payload.status not in allowed:
        raise HTTPException(400, "Status de entrega invalido")
    manifest = get_delivery_manifest(db, manifest_id)
    if not has_permission(user, "manifests.create") and manifest.driver_name.lower() != user.name.lower():
        raise HTTPException(403, "Romaneio atribuido a outro motorista")
    if manifest.status == "cancelado":
        raise HTTPException(400, "Romaneio cancelado")
    item = next((row for row in manifest.items if row.id == item_id), None)
    if not item:
        raise HTTPException(404, "Entrega nao encontrada")
    item.status = payload.status
    item.note = payload.note
    item.delivered_at = datetime.utcnow() if payload.status == "entregue" else None
    if manifest.status == "preparacao":
        manifest.status = "em_rota"
    if all(row.status == "entregue" for row in manifest.items):
        manifest.status = "concluido"
    audit(db, user, "status", "delivery_manifest_item", item.id, payload.status)
    db.commit()
    return delivery_manifest_to_schema(get_delivery_manifest(db, manifest.id))


@app.get("/api/delivery-manifests/{manifest_id}/pdf")
def delivery_manifest_pdf(manifest_id: int, db: Session = Depends(get_db), user: models.User = Depends(require_permission("manifests.print"))):
    manifest = get_delivery_manifest(db, manifest_id)
    stream = BytesIO()
    doc = SimpleDocTemplate(
        stream,
        pagesize=A4,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"Romaneio {manifest.code}",
        author=settings.app_name,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ManifestTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=colors.HexColor("#17201a"), spaceAfter=3 * mm)
    subtitle_style = ParagraphStyle("ManifestSubtitle", parent=styles["Normal"], fontSize=9, leading=12, textColor=colors.HexColor("#66756b"))
    cell_style = ParagraphStyle("ManifestCell", parent=styles["Normal"], fontSize=8, leading=10, textColor=colors.HexColor("#17201a"))
    cell_bold_style = ParagraphStyle("ManifestCellBold", parent=cell_style, fontName="Helvetica-Bold")
    small_style = ParagraphStyle("ManifestSmall", parent=styles["Normal"], fontSize=7.5, leading=9, textColor=colors.HexColor("#66756b"))
    status_labels = {"preparacao": "EM PREPARACAO", "em_rota": "EM ROTA", "concluido": "CONCLUIDO", "cancelado": "CANCELADO", "pendente": "PENDENTE", "entregue": "ENTREGUE", "nao_entregue": "NAO ENTREGUE"}

    story = [
        Paragraph("CARVAO PRO", subtitle_style),
        Paragraph("Romaneio de Entregas", title_style),
        Paragraph(manifest.code, subtitle_style),
        Spacer(1, 2 * mm),
    ]
    details = [
        [Paragraph("DATA DA ENTREGA", small_style), Paragraph("MOTORISTA", small_style), Paragraph("VEICULO", small_style), Paragraph("STATUS", small_style)],
        [Paragraph(manifest.delivery_date.strftime("%d/%m/%Y"), cell_bold_style), Paragraph(escape(manifest.driver_name), cell_bold_style), Paragraph(escape(manifest.vehicle or "Nao informado"), cell_style), Paragraph(status_labels.get(manifest.status, manifest.status.upper()), cell_bold_style)],
    ]
    details_table = Table(details, colWidths=[34 * mm, 58 * mm, 48 * mm, 37 * mm])
    details_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef4ed")),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#cfd8cd")),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dfe6dd")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.extend([details_table, Spacer(1, 4 * mm)])
    if manifest.notes:
        story.extend([Paragraph(f"<b>Observacoes:</b> {escape(manifest.notes)}", cell_style), Spacer(1, 4 * mm)])

    rows = [[
        Paragraph("ORDEM", small_style),
        Paragraph("CLIENTE / CONTATO", small_style),
        Paragraph("ENDERECO", small_style),
        Paragraph("PRODUTOS", small_style),
        Paragraph("VALOR", small_style),
        Paragraph("ENTREGA", small_style),
    ]]
    total_value = 0.0
    for item in manifest.items:
        sale = item.sale
        products = "<br/>".join(f"{sale_item.quantity:g}x {escape(sale_item.product.name)}" for sale_item in sale.items)
        total_value += sale.total_value
        rows.append([
            Paragraph(str(item.delivery_order), cell_bold_style),
            Paragraph(f"<b>{escape(sale.customer_name or f'Venda #{sale.id}')}</b><br/>{escape(sale.customer_phone or 'Sem telefone')}", cell_style),
            Paragraph(escape(item.delivery_address), cell_style),
            Paragraph(products, cell_style),
            Paragraph(f"R$ {sale.total_value:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."), cell_style),
            Paragraph(status_labels.get(item.status, item.status.upper()), cell_bold_style),
        ])
    deliveries_table = Table(rows, colWidths=[17 * mm, 37 * mm, 45 * mm, 35 * mm, 24 * mm, 25 * mm], repeatRows=1)
    deliveries_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#18231c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cfd8cd")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f9f6")]),
    ]))
    total_label = f"{len(manifest.items)} entrega(s) - Total R$ {total_value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    story.extend([deliveries_table, Spacer(1, 4 * mm), Paragraph(f"<b>{total_label}</b>", cell_style), Spacer(1, 14 * mm)])
    signatures = Table([
        ["________________________________", "________________________________"],
        ["Motorista", "Responsavel pela expedicao"],
    ], colWidths=[85 * mm, 85 * mm])
    signatures.setStyle(TableStyle([
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica"),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#66756b")),
    ]))
    story.append(signatures)

    def draw_page_number(pdf_canvas, pdf_doc):
        pdf_canvas.saveState()
        pdf_canvas.setFont("Helvetica", 8)
        pdf_canvas.setFillColor(colors.HexColor("#66756b"))
        pdf_canvas.drawString(14 * mm, 9 * mm, f"Emitido em {datetime.now():%d/%m/%Y %H:%M}")
        pdf_canvas.drawRightString(A4[0] - 14 * mm, 9 * mm, f"Pagina {pdf_doc.page}")
        pdf_canvas.restoreState()

    doc.build(story, onFirstPage=draw_page_number, onLaterPages=draw_page_number)
    filename = f"romaneio-{manifest.code}.pdf"
    return Response(stream.getvalue(), media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="{filename}"'})


@app.get("/api/finance/dashboard")
def finance_dashboard(db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.view_receivables", "finance.view_payables", match_all=False))):
    today = date.today()
    receivables = db.query(models.Receivable).filter(models.Receivable.status != "cancelado").all()
    payables = db.query(models.Payable).filter(models.Payable.status != "cancelado").all()
    accounts = db.query(models.FinancialAccount).filter(models.FinancialAccount.active.is_(True)).all()
    received_today = db.query(func.coalesce(func.sum(models.CashTransaction.amount), 0)).filter(func.date(models.CashTransaction.occurred_at) == today, models.CashTransaction.direction == "entrada", models.CashTransaction.reversed.is_(False)).scalar() or 0
    paid_today = db.query(func.coalesce(func.sum(models.CashTransaction.amount), 0)).filter(func.date(models.CashTransaction.occurred_at) == today, models.CashTransaction.direction == "saida", models.CashTransaction.reversed.is_(False)).scalar() or 0
    return {
        "cash_balance": sum(financial_account_balance(db, account) for account in accounts),
        "receivable_open": sum(max(0, row.original_amount + row.interest_amount + row.fine_amount - row.discount_amount - row.paid_amount) for row in receivables),
        "payable_open": sum(max(0, row.original_amount - row.paid_amount) for row in payables),
        "received_today": float(received_today), "paid_today": float(paid_today),
        "overdue_customers": len({row.customer_id for row in receivables if row.due_date < today and row.paid_amount < row.original_amount}),
    }


@app.get("/api/finance/accounts")
def list_financial_accounts(db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.view_cash_flow"))):
    accounts = db.query(models.FinancialAccount).order_by(models.FinancialAccount.name).all()
    return [{"id": row.id, "name": row.name, "account_type": row.account_type, "initial_balance": row.initial_balance, "active": row.active, "balance": financial_account_balance(db, row)} for row in accounts]


@app.post("/api/finance/accounts")
def create_financial_account(payload: schemas.FinancialAccountCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.manage_accounts"))):
    if db.query(models.FinancialAccount).filter(func.lower(models.FinancialAccount.name) == payload.name.lower()).first(): raise HTTPException(400, "Conta financeira ja cadastrada")
    row = models.FinancialAccount(**payload.model_dump()); db.add(row); db.flush(); audit(db, user, "create", "financial_account", row.id, row.name); db.commit()
    return {"id": row.id, **payload.model_dump(), "active": True, "balance": row.initial_balance}


@app.get("/api/finance/cost-centers", response_model=list[schemas.CostCenterOut])
def list_cost_centers(db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.view_payables"))):
    return db.query(models.CostCenter).filter(models.CostCenter.active.is_(True)).order_by(models.CostCenter.name).all()


@app.post("/api/finance/cost-centers", response_model=schemas.CostCenterOut)
def create_cost_center(payload: schemas.NamedEntityCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("settings.financial"))):
    row = models.CostCenter(name=payload.name.strip()); db.add(row); db.flush(); audit(db, user, "create", "cost_center", row.id, row.name); db.commit(); db.refresh(row); return row


@app.get("/api/finance/suppliers", response_model=list[schemas.SupplierOut])
def list_suppliers(db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.view_payables"))):
    return db.query(models.Supplier).filter(models.Supplier.active.is_(True)).order_by(models.Supplier.name).all()


@app.post("/api/finance/suppliers", response_model=schemas.SupplierOut)
def create_supplier(payload: schemas.SupplierCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.create_payables"))):
    row = models.Supplier(**payload.model_dump()); db.add(row); db.flush(); audit(db, user, "create", "supplier", row.id, row.name); db.commit(); db.refresh(row); return row


@app.get("/api/finance/receivables")
def list_receivables(db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.view_receivables"))):
    rows = db.query(models.Receivable).options(joinedload(models.Receivable.customer), joinedload(models.Receivable.seller)).order_by(models.Receivable.due_date).all()
    today = date.today()
    return [{"id": row.id, "sale_id": row.sale_id, "customer_name": row.customer.legal_name, "seller_name": row.seller.name, "due_date": row.due_date, "original_amount": row.original_amount, "paid_amount": row.paid_amount, "balance": max(0, row.original_amount + row.interest_amount + row.fine_amount - row.discount_amount - row.paid_amount), "status": "vencido" if row.status in {"aberto", "parcial"} and row.due_date < today else row.status} for row in rows]


@app.post("/api/finance/receivables/{receivable_id}/payments")
def receive_payment(receivable_id: int, payload: schemas.FinancialPaymentCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.settle_titles"))):
    row = db.get(models.Receivable, receivable_id); account = db.get(models.FinancialAccount, payload.account_id)
    if not row or row.status == "cancelado": raise HTTPException(404, "Titulo a receber nao encontrado")
    if not account or not account.active: raise HTTPException(404, "Conta financeira nao encontrada")
    balance = row.original_amount + row.interest_amount + row.fine_amount - row.discount_amount - row.paid_amount
    if payload.amount > balance + 0.001: raise HTTPException(400, "Valor maior que o saldo em aberto")
    payment = models.ReceivablePayment(receivable_id=row.id, account_id=account.id, amount=payload.amount, payment_method=payload.payment_method, occurred_at=payload.occurred_at or datetime.utcnow(), note=payload.note, responsible_id=user.id)
    db.add(payment); db.flush(); row.paid_amount += payload.amount; row.status = "pago" if row.paid_amount >= row.original_amount + row.interest_amount + row.fine_amount - row.discount_amount else "parcial"
    db.add(models.CashTransaction(account_id=account.id, direction="entrada", amount=payload.amount, occurred_at=payment.occurred_at, source_type="receivable_payment", source_id=payment.id, description=f"Recebimento do titulo #{row.id}", responsible_id=user.id))
    audit(db, user, "receive", "receivable", row.id, f"Baixa R$ {payload.amount:.2f}"); db.commit(); return {"ok": True, "status": row.status, "paid_amount": row.paid_amount}


@app.get("/api/finance/payables")
def list_payables(db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.view_payables"))):
    rows = db.query(models.Payable).options(joinedload(models.Payable.supplier), joinedload(models.Payable.cost_center)).order_by(models.Payable.due_date).all(); today = date.today()
    return [{"id": row.id, "description": row.description, "supplier_name": row.supplier.name if row.supplier else None, "cost_center_name": row.cost_center.name, "category": row.category, "due_date": row.due_date, "original_amount": row.original_amount, "paid_amount": row.paid_amount, "balance": max(0, row.original_amount - row.paid_amount), "status": "vencido" if row.status in {"aberto", "parcial"} and row.due_date < today else row.status} for row in rows]


@app.post("/api/finance/payables")
def create_payable(payload: schemas.PayableCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.create_payables"))):
    if not db.get(models.CostCenter, payload.cost_center_id): raise HTTPException(404, "Centro de custo nao encontrado")
    if payload.supplier_id and not db.get(models.Supplier, payload.supplier_id): raise HTTPException(404, "Fornecedor nao encontrado")
    row = models.Payable(**payload.model_dump()); db.add(row); db.flush(); audit(db, user, "create", "payable", row.id, f"{row.description} R$ {row.original_amount:.2f}"); db.commit(); return {"id": row.id, "status": row.status}


@app.post("/api/finance/payables/{payable_id}/payments")
def pay_payable(payable_id: int, payload: schemas.FinancialPaymentCreate, db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.settle_titles"))):
    row = db.get(models.Payable, payable_id); account = db.get(models.FinancialAccount, payload.account_id)
    if not row or row.status == "cancelado": raise HTTPException(404, "Conta a pagar nao encontrada")
    if not account or not account.active: raise HTTPException(404, "Conta financeira nao encontrada")
    if payload.amount > row.original_amount - row.paid_amount + 0.001: raise HTTPException(400, "Valor maior que o saldo em aberto")
    payment = models.PayablePayment(payable_id=row.id, account_id=account.id, amount=payload.amount, occurred_at=payload.occurred_at or datetime.utcnow(), note=payload.note, responsible_id=user.id)
    db.add(payment); db.flush(); row.paid_amount += payload.amount; row.status = "pago" if row.paid_amount >= row.original_amount else "parcial"
    db.add(models.CashTransaction(account_id=account.id, direction="saida", amount=payload.amount, occurred_at=payment.occurred_at, source_type="payable_payment", source_id=payment.id, description=f"Pagamento da conta #{row.id}: {row.description}", responsible_id=user.id))
    audit(db, user, "pay", "payable", row.id, f"Pagamento R$ {payload.amount:.2f}"); db.commit(); return {"ok": True, "status": row.status, "paid_amount": row.paid_amount}


@app.get("/api/finance/cash-flow")
def cash_flow(start: date | None = None, end: date | None = None, db: Session = Depends(get_db), user: models.User = Depends(require_permission("finance.view_cash_flow"))):
    query = db.query(models.CashTransaction).options(joinedload(models.CashTransaction.account)).filter(models.CashTransaction.reversed.is_(False))
    if start: query = query.filter(models.CashTransaction.occurred_at >= datetime.combine(start, time.min))
    if end: query = query.filter(models.CashTransaction.occurred_at <= datetime.combine(end, time.max))
    rows = query.order_by(models.CashTransaction.occurred_at.desc()).limit(500).all()
    return [{"id": row.id, "account_name": row.account.name, "direction": row.direction, "amount": row.amount, "occurred_at": row.occurred_at, "description": row.description} for row in rows]


@app.get("/api/dashboard")
def dashboard(
    start: date | None = Query(None),
    end: date | None = Query(None),
    seller_id: int | None = None,
    db: Session = Depends(get_db),
    user: models.User = Depends(require_permission("dashboard.view")),
):
    today = date.today()
    start_dt = datetime.combine(start or today, time.min)
    end_dt = datetime.combine(end or today, time.max)
    query = db.query(models.Sale).filter(models.Sale.status == "confirmada", models.Sale.occurred_at >= start_dt, models.Sale.occurred_at <= end_dt)
    if has_permission(user, "dashboard.view_own") and not has_permission(user, "dashboard.view_general"):
        query = query.filter(models.Sale.seller_id == user.seller_id)
    elif seller_id:
        query = query.filter(models.Sale.seller_id == seller_id)

    sales = query.all()
    total_value = sum(s.total_value for s in sales)
    sale_ids = [s.id for s in sales] or [0]
    item_qty = db.query(func.coalesce(func.sum(models.SaleItem.quantity), 0)).filter(models.SaleItem.sale_id.in_(sale_ids)).scalar()
    sellers = db.query(models.Seller).all()
    ranking = []
    for seller in sellers:
        seller_sales = [s for s in sales if s.seller_id == seller.id]
        qty = db.query(func.coalesce(func.sum(models.SaleItem.quantity), 0)).join(models.Sale).filter(models.Sale.seller_id == seller.id, models.Sale.id.in_(sale_ids)).scalar()
        value = sum(s.total_value for s in seller_sales)
        ranking.append({"seller_id": seller.id, "seller_name": seller.name, "quantity": float(qty or 0), "value": value, "commission": value * ((seller.commission_percent or 0) / 100)})
    ranking.sort(key=lambda row: row["value"], reverse=True)
    products = db.query(models.Product).order_by(models.Product.type, models.Product.name).all()
    return {
        "period": {"start": start_dt.date(), "end": end_dt.date()},
        "sales_count": len(sales),
        "total_value": total_value,
        "items_quantity": float(item_qty or 0),
        "stock": [{"id": p.id, "name": p.name, "type": p.type, "unit": p.unit, "current_stock": p.current_stock, "minimum_stock": p.minimum_stock, "low": p.current_stock <= p.minimum_stock} for p in products],
        "ranking": ranking,
    }


@app.get("/api/whatsapp/settings", response_model=schemas.WhatsAppSettingsOut)
def read_whatsapp(db: Session = Depends(get_db), user: models.User = Depends(require_permission("settings.alerts"))):
    return get_whatsapp_settings(db)


@app.put("/api/whatsapp/settings", response_model=schemas.WhatsAppSettingsOut)
def update_whatsapp(payload: schemas.WhatsAppSettingsIn, db: Session = Depends(get_db), user: models.User = Depends(require_permission("settings.alerts"))):
    settings_row = get_whatsapp_settings(db)
    for key, value in payload.model_dump().items():
        setattr(settings_row, key, value)
    audit(db, user, "update", "whatsapp_settings", settings_row.id, "Configuracao atualizada")
    db.commit()
    db.refresh(settings_row)
    return settings_row


@app.post("/api/whatsapp/test")
async def test_whatsapp(db: Session = Depends(get_db), user: models.User = Depends(require_permission("settings.integrations"))):
    settings_row = get_whatsapp_settings(db)
    return await send_whatsapp(db, settings_row.manager_phone, "Teste de integracao WhatsApp do Sistema Carvao.")


@app.get("/api/reports/export.xlsx")
def export_xlsx(db: Session = Depends(get_db), user: models.User = Depends(require_permission("reports.export_excel"))):
    sales = db.query(models.Sale).options(joinedload(models.Sale.seller)).order_by(models.Sale.occurred_at.desc()).all()
    sellers = db.query(models.Seller).order_by(models.Seller.name).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    headers = ["ID", "Data", "Vendedor", "Pagamento", "Status", "Total", "Comissao (%)", "Comissao (R$)"]
    widths = {"A": 10, "B": 20, "C": 24, "D": 16, "E": 14, "F": 16, "G": 15, "H": 18}

    def fill_sales_sheet(sheet, rows, include_totals=False):
        sheet.append(headers)
        for sale in rows:
            percent = sale.seller.commission_percent or 0
            commission = sale.total_value * (percent / 100) if sale.status == "confirmada" else 0
            sheet.append([sale.id, sale.occurred_at.strftime("%d/%m/%Y %H:%M"), sale.seller.name, sale.payment_method, sale.status, sale.total_value, percent, commission])
        for cell in sheet[1]:
            cell.font = cell.font.copy(bold=True)
        sheet.freeze_panes = "A2"
        data_end = max(sheet.max_row, 1)
        sheet.auto_filter.ref = f"A1:H{data_end}"
        for row_number in range(2, data_end + 1):
            sheet.cell(row_number, 6).number_format = 'R$ #,##0.00'
            sheet.cell(row_number, 7).number_format = '0.00"%"'
            sheet.cell(row_number, 8).number_format = 'R$ #,##0.00'
        if include_totals:
            confirmed = [sale for sale in rows if sale.status == "confirmada"]
            total_sold = sum(sale.total_value for sale in confirmed)
            total_commission = sum(sale.total_value * ((sale.seller.commission_percent or 0) / 100) for sale in confirmed)
            sheet.append([])
            sheet.append(["TOTAIS", "", "", "", f"{len(confirmed)} vendas confirmadas", total_sold, "", total_commission])
            total_row = sheet.max_row
            sheet.cell(total_row, 1).font = sheet.cell(total_row, 1).font.copy(bold=True)
            sheet.cell(total_row, 5).font = sheet.cell(total_row, 5).font.copy(bold=True)
            sheet.cell(total_row, 6).font = sheet.cell(total_row, 6).font.copy(bold=True)
            sheet.cell(total_row, 8).font = sheet.cell(total_row, 8).font.copy(bold=True)
            sheet.cell(total_row, 6).number_format = 'R$ #,##0.00'
            sheet.cell(total_row, 8).number_format = 'R$ #,##0.00'
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

    fill_sales_sheet(ws, sales)

    summary = wb.create_sheet("Comissoes por vendedor")
    summary.append(["Vendedor", "Vendas confirmadas", "Total vendido", "Comissao (%)", "Comissao total"])
    for cell in summary[1]:
        cell.font = cell.font.copy(bold=True)
    for seller in sellers:
        seller_sales = [sale for sale in sales if sale.seller_id == seller.id and sale.status == "confirmada"]
        sold = sum(sale.total_value for sale in seller_sales)
        percent = seller.commission_percent or 0
        summary.append([seller.name, len(seller_sales), sold, percent, sold * (percent / 100)])
    for row in range(2, summary.max_row + 1):
        summary.cell(row, 3).number_format = 'R$ #,##0.00'
        summary.cell(row, 4).number_format = '0.00"%"'
        summary.cell(row, 5).number_format = 'R$ #,##0.00'
    for column, width in {"A": 26, "B": 21, "C": 18, "D": 15, "E": 18}.items():
        summary.column_dimensions[column].width = width

    used_titles = set(wb.sheetnames)
    for seller in sellers:
        base_title = "".join("-" if character in '[]:*?/\\' else character for character in seller.name).strip() or f"Vendedor {seller.id}"
        base_title = base_title[:31]
        title = base_title
        suffix = 2
        while title in used_titles:
            marker = f" ({suffix})"
            title = f"{base_title[:31 - len(marker)]}{marker}"
            suffix += 1
        used_titles.add(title)
        seller_sheet = wb.create_sheet(title)
        fill_sales_sheet(seller_sheet, [sale for sale in sales if sale.seller_id == seller.id], include_totals=True)
    stream = BytesIO()
    wb.save(stream)
    return Response(stream.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=relatorio-vendas.xlsx"})


@app.get("/api/reports/export.pdf")
def export_pdf(db: Session = Depends(get_db), user: models.User = Depends(require_permission("reports.export_pdf"))):
    sales = db.query(models.Sale).options(joinedload(models.Sale.seller)).order_by(models.Sale.occurred_at.desc()).all()
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm)
    styles = getSampleStyleSheet()
    story = [Paragraph("Relatorio de Vendas - Sistema Carvao", styles["Title"]), Spacer(1, 5 * mm)]
    rows = [["ID", "Data", "Vendedor", "Pagamento", "Status", "Total", "%", "Comissao"]]
    for sale in sales:
        percent = sale.seller.commission_percent or 0
        commission = sale.total_value * (percent / 100) if sale.status == "confirmada" else 0
        rows.append([str(sale.id), sale.occurred_at.strftime("%d/%m/%Y %H:%M"), sale.seller.name, sale.payment_method, sale.status, f"R$ {sale.total_value:.2f}", f"{percent:.2f}%", f"R$ {commission:.2f}"])
    sales_table = Table(rows, repeatRows=1, colWidths=[14 * mm, 34 * mm, 43 * mm, 27 * mm, 24 * mm, 28 * mm, 18 * mm, 30 * mm])
    sales_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#18231c")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cfd8cd")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f1")]),
        ("ALIGN", (5, 1), (-1, -1), "RIGHT"),
    ]))
    story.extend([sales_table, Spacer(1, 7 * mm), Paragraph("Resumo de comissoes por vendedor", styles["Heading2"])])
    summary_rows = [["Vendedor", "Vendas confirmadas", "Total vendido", "%", "Comissao total"]]
    for seller in db.query(models.Seller).order_by(models.Seller.name).all():
        seller_sales = [sale for sale in sales if sale.seller_id == seller.id and sale.status == "confirmada"]
        sold = sum(sale.total_value for sale in seller_sales)
        percent = seller.commission_percent or 0
        summary_rows.append([seller.name, str(len(seller_sales)), f"R$ {sold:.2f}", f"{percent:.2f}%", f"R$ {sold * (percent / 100):.2f}"])
    summary_table = Table(summary_rows, repeatRows=1, colWidths=[55 * mm, 38 * mm, 38 * mm, 22 * mm, 40 * mm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ffbf47")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cfd8cd")),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
    ]))
    story.append(summary_table)

    def report_footer(pdf_canvas, pdf_doc):
        pdf_canvas.saveState(); pdf_canvas.setFont("Helvetica", 8)
        pdf_canvas.drawString(12 * mm, 7 * mm, f"Emitido em {datetime.now():%d/%m/%Y %H:%M}")
        pdf_canvas.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"Pagina {pdf_doc.page}")
        pdf_canvas.restoreState()

    doc.build(story, onFirstPage=report_footer, onLaterPages=report_footer)
    return Response(stream.getvalue(), media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=relatorio-vendas.pdf"})


@app.post("/api/whatsapp/webhook")
async def whatsapp_webhook(payload: dict):
    return {"received": True, "phase": "parser-ready", "payload": payload}
